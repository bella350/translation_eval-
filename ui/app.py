import sys
import json
from pathlib import Path
from datetime import datetime

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import streamlit as st
import plotly.graph_objects as go
from docx import Document
from core.scorer import evaluate

# ================= 历史持久化 =================
HISTORY_FILE = Path(__file__).resolve().parent.parent / "history.json"

def _safe_for_json(obj):
    if isinstance(obj, dict):
        return {k: _safe_for_json(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_safe_for_json(v) for v in obj]
    if hasattr(obj, "item"):
        return obj.item()
    if hasattr(obj, "tolist"):
        return obj.tolist()
    return obj

def load_history():
    if HISTORY_FILE.exists():
        try:
            return json.loads(HISTORY_FILE.read_text(encoding="utf-8"))
        except Exception:
            return []
    return []

def save_history(hist):
    HISTORY_FILE.write_text(
        json.dumps(_safe_for_json(hist), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
# =============================================

st.set_page_config(page_title="翻译评估工具", layout="wide")
st.sidebar.title("导航")

PAGES = ["上传与评估", "图表", "导出", "历史与对比"]
tab = st.sidebar.radio(
    "选择页面",
    PAGES,
    index=PAGES.index(st.session_state.get("_switch_tab", "上传与评估")),
)
if "_switch_tab" in st.session_state:
    del st.session_state["_switch_tab"]

if "history" not in st.session_state:
    st.session_state["history"] = load_history()

# ============== 页面 1：上传与评估 ==============
if tab == "上传与评估":
    st.header("上传与评估")

    # 回填：有数据就读取，读完立即删除，避免反复跳回
    prefill = st.session_state.pop("prefill", None)

    src = st.text_area(
        "原文",
        value=prefill["src"] if prefill else st.session_state.get("src_input", ""),
        height=150,
        key="src_input",
    )
    tgt = st.text_area(
        "译文",
        value=prefill["tgt"] if prefill else st.session_state.get("tgt_input", ""),
        height=150,
        key="tgt_input",
    )
    glossary_raw = st.text_area(
        "术语表（可选，一行一对，格式：源术语 → 目标术语）",
        value=(
            "\n".join(f"{k} → {v}" for k, v in prefill["glossary"].items())
            if prefill and prefill.get("glossary")
            else st.session_state.get("glossary_input", "")
        ),
        placeholder="创新点 → инновационные моменты\n理论 → теоретический\n方法 → методический\n应用 → прикладной",
        height=100,
        key="glossary_input",
    )
    ref = st.text_area(
        "参考译文（可选，用于 BLEU/chrF 对比）",
        value=prefill["ref"] if prefill else st.session_state.get("ref_input", ""),
        height=120,
        placeholder="留空则仅做现有5维评估；填写后展示参考对齐指标",
        key="ref_input",
    )

    # 解析术语表
    glossary = {}
    if glossary_raw.strip():
        for line in glossary_raw.strip().split("\n"):
            line = line.strip()
            if "→" in line:
                parts = line.split("→", 1)
                s, t = parts[0].strip(), parts[1].strip()
                if s and t:
                    glossary[s] = t

    if st.button("评估", key="eval_btn"):
        # ★ 加载提示（防首次调用 5~15 秒误判"没反应"）
        with st.spinner("⏳ 正在调用通义千问评估中，首次约 5~15 秒，请稍候…"):
            result = evaluate(src, tgt, glossary=glossary, ref=ref if ref.strip() else None)

        # ★ 模式标识 + 总评摘要
        mode = result.get("mode", "")
        if "llm" in mode:
            st.success(f"✅ 综合评分：{result['final_score']:.1f} ｜ 评判模式：LLM 大模型评分（通义千问）")
            if result.get("summary"):
                st.info("📋 总评：" + result["summary"])
        else:
            st.warning(f"⚠️ 综合评分：{result['final_score']:.1f} ｜ 评判模式：规则兜底（LLM 未响应/超时）")

        # 维度明细（原样保留，展开显示）
        with st.expander("维度明细", expanded=True):
            for dim in result["dimensions"]:
                st.write(f"**{dim['dimension']}** ({dim['score']}/{dim['weight']})")
                for note in dim.get("notes", []):
                    st.write(f"- {note}")

        st.session_state["score"] = float(result["final_score"])
        st.session_state["eval_details"] = result["dimensions"]

        # 写入历史（持久化）
        st.session_state["history"].append({
            "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "src": src,
            "tgt": tgt,
            "glossary": glossary,
            "ref": ref if ref.strip() else None,
            "score": float(result["final_score"]),
            "dimensions": result["dimensions"],
            "mode": mode,
            "summary": result.get("summary", ""),
        })
        st.session_state["history"] = st.session_state["history"][-50:]
        save_history(st.session_state["history"])
        st.success(f"已记录到历史（共 {len(st.session_state['history'])} 条）")

# ============== 页面 2：图表 ==============
elif tab == "图表":
    st.header("评分维度雷达图")
    if "eval_details" in st.session_state:
        dims = [d for d in st.session_state["eval_details"] if d["weight"] > 0]
        vals = [d["score"] for d in dims]
        theta = [d["dimension"] for d in dims]
        fig = go.Figure(data=go.Scatterpolar(
            r=vals + vals[:1],
            theta=theta + theta[:1],
            fill="toself",
        ))
        fig.update_layout(
            polar=dict(radialaxis=dict(range=[0, max(d["weight"] for d in dims)])),
            showlegend=False,
        )
        st.plotly_chart(fig, use_container_width=True)
    else:
        st.info("请先在「上传与评估」页面完成评估")

# ============== 页面 3：导出 ==============
elif tab == "导出":
    st.header("导出 Word 报告")
    if st.button("生成 docx"):
        d = Document()
        d.add_heading("翻译评估结果", 0)
        d.add_paragraph(f"综合评分：{st.session_state.get('score', '未评估')}")
        if "eval_details" in st.session_state:
            d.add_heading("维度明细", level=1)
            for dim in st.session_state["eval_details"]:
                d.add_paragraph(f"{dim['dimension']}：{dim['score']}/{dim['weight']}")
                for note in dim.get("notes", []):
                    d.add_paragraph(f"  - {note}", style="List Bullet")
        d.save("result.docx")
        st.success("已生成 result.docx（在命令执行目录下）")

# ============== 页面 4：历史与对比 ==============
else:
    st.header("历史记录与对比")
    hist = st.session_state.get("history", [])

    c1, c2 = st.columns([1, 1])
    with c1:
        if st.button("清空历史"):
            st.session_state["history"] = []
            save_history([])
            st.experimental_rerun()
    with c2:
        st.write(f"共 {len(hist)} 条（保留最近 50 条）")

    if not hist:
        st.info("暂无评估记录，先去「上传与评估」跑一次")
    else:
        # 概览表格
        rows = []
        for i, h in enumerate(hist):
            rows.append({
                "#": i,
                "时间": h["time"],
                "得分": f"{h['score']:.1f}",
                "模式": h.get("mode", "未知"),
                "原文(截断)": (h.get("src") or "")[:40].replace("\n", " "),
                "译文(截断)": (h.get("tgt") or "")[:40].replace("\n", " "),
            })
        st.dataframe(rows, use_container_width=True)
        st.divider()

        # 单条详情 + 回填按钮
        for i, h in enumerate(reversed(hist)):
            real_idx = len(hist) - 1 - i
            with st.expander(f"#{real_idx}  {h['time']}  |  得分 {h['score']:.1f}  |  {h.get('mode', '')}"):

                # ★ 回填按钮
                if st.button("⬇ 回填到评估页并重测", key=f"prefill_{real_idx}"):
                    st.session_state["prefill"] = {
                        "src": h.get("src", ""),
                        "tgt": h.get("tgt", ""),
                        "glossary": h.get("glossary", {}),
                        "ref": h.get("ref", "") or "",
                    }
                    st.session_state["_switch_tab"] = "上传与评估"
                    st.experimental_rerun()

                if h.get("summary"):
                    st.info("📋 " + h["summary"])

                st.write("**原文**")
                st.text_area("", h.get("src", ""), height=80, key=f"h_src_{real_idx}")
                st.write("**译文**")
                st.text_area("", h.get("tgt", ""), height=80, key=f"h_tgt_{real_idx}")
                if h.get("glossary"):
                    st.write("**术语表**")
                    st.code("\n".join(f"{k} → {v}" for k, v in h["glossary"].items()))
                st.write("**维度明细**")
                for dim in h["dimensions"]:
                    st.write(f"- **{dim['dimension']}** ({dim['score']}/{dim['weight']})")
                    for note in dim.get("notes", []):
                        st.write(f"  - {note}")

            st.divider()

        # 对比两条
        st.subheader("对比两条记录")
        opts = {f"#{i}  {hist[i]['time']}  得分{hist[i]['score']:.1f}": i
                for i in range(len(hist))}
        sel_a = st.selectbox("记录 A", list(opts.keys()), index=0, key="sel_a")
        sel_b = st.selectbox("记录 B", list(opts.keys()), index=len(opts) - 1, key="sel_b")

        if st.button("对比 A vs B"):
            ia, ib = opts[sel_a], opts[sel_b]
            ha, hb = hist[ia], hist[ib]

            cc1, cc2 = st.columns(2)
            with cc1:
                st.markdown(f"**A #{ia}**｜{ha['time']}｜得分 **{ha['score']:.1f}**")
                st.text_area("译文 A", ha.get("tgt", ""), height=140, key="cmp_a")
                if ha.get("ref"):
                    st.caption(f"参考译文 A：{ha['ref'][:60]}")
            with cc2:
                st.markdown(f"**B #{ib}**｜{hb['time']}｜得分 **{hb['score']:.1f}**")
                st.text_area("译文 B", hb.get("tgt", ""), height=140, key="cmp_b")
                if hb.get("ref"):
                    st.caption(f"参考译文 B：{hb['ref'][:60]}")

            st.markdown("**维度差异**")
            da = {d["dimension"]: d for d in ha["dimensions"]}
            db = {d["dimension"]: d for d in hb["dimensions"]}
            all_keys = list(dict.fromkeys(list(da.keys()) + list(db.keys())))
            for k in all_keys:
                sa = da[k]["score"] if k in da else 0
                sb = db[k]["score"] if k in db else 0
                diff = sa - sb
                if diff > 0:
                    tag = f"🟢 A 高 {diff}"
                elif diff < 0:
                    tag = f"🔴 B 高 {abs(diff)}"
                else:
                    tag = "⚪ 持平"
                st.write(f"- {k}：A={sa}  B={sb}　**{tag}**")

            # 若有参考译文，额外对比 BLEU/chrF
            for name, rec in (("A", ha), ("B", hb)):
                for d in rec["dimensions"]:
                    if d["dimension"] == "参考对齐(BLEU/chrF)":
                        bleu = d.get("bleu")
                        chrf = d.get("chrf")
                        if bleu is not None:
                            st.caption(f"{name} 的参考指标：BLEU={bleu:.2f}  chrF={chrf:.2f}")